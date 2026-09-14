"""
One-time import of Justin's baseline tracking workbook into analyst records.

Rules (agreed 2026-09-10):
  * Notes are copied VERBATIM, prefixed "[Justin workbook <sheet>] ".
  * A status is only *suggested* from literal text in the note:
        "Moratorium until ..."  -> Existing Moratorium
        note starts "Monitor"   -> Monitor
        "Moratorium expired"    -> No Action
    and is marked status_source = "baseline <sheet> (unconfirmed)" until the
    analyst saves or confirms the record in the app.
  * Never overwrite an existing analyst note or status. Idempotent.
  * Justin's Distance column is kept as `baseline_reference` (it is the WFIGS
    origin-point miles he sorted by, not perimeter distance).
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

import analyst

BASE_DIR = Path(__file__).resolve().parent
REFERENCE_DIR = BASE_DIR / "reference"
DEFAULT_WORKBOOK = REFERENCE_DIR / "2026-09 Fire Tracking.xlsx"
DEFAULT_SHEET = "9-8"        # sheet carrying Justin's written conclusions
REVIEW_SHEET = "9-9"         # last completed review pass (no new notes = position unchanged)

_RULES = [
    (re.compile(r"\bmoratorium\s+until\b", re.I), "Existing Moratorium"),
    (re.compile(r"^\s*monitor\b", re.I), "Monitor"),
    (re.compile(r"\bmoratorium\s+expired\b", re.I), "No Action"),
]


def suggest_status(note: str) -> str | None:
    """Literal keyword match only. None when nothing matches."""
    for rx, status in _RULES:
        if rx.search(note or ""):
            return status
    return None


def find_workbooks() -> list[Path]:
    if not REFERENCE_DIR.exists():
        return []
    return sorted(REFERENCE_DIR.glob("*.xlsx"))


def sheet_names(workbook: Path) -> list[str]:
    wb = openpyxl.load_workbook(workbook, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def read_sheet(workbook: Path, sheet: str) -> list[dict[str, Any]]:
    """Rows of the sheet that carry a note, as dicts of the columns we use."""
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]

    def col(name: str) -> int | None:
        return header.index(name) if name in header else None

    c_irwin, c_name, c_notes = col("attr_IrwinID"), col("poly_IncidentName"), col("Notes")
    c_dist = col("Distance") if col("Distance") is not None else col("Miles")
    c_loc = col("Location")
    if c_irwin is None or c_notes is None:
        raise ValueError(f"Sheet {sheet!r} lacks attr_IrwinID / Notes columns")

    out = []
    for r in rows[1:]:
        note = r[c_notes] if c_notes < len(r) else None
        iid = r[c_irwin] if c_irwin < len(r) else None
        if not iid or note in (None, ""):
            continue
        out.append({
            "irwin_id": str(iid).strip().upper(),
            "fire_name": r[c_name] if c_name is not None else None,
            "note": str(note).strip(),
            "distance": r[c_dist] if c_dist is not None else None,
            "location": r[c_loc] if c_loc is not None else None,
        })
    return out


# --------------------------------------------------------------------------- #
# Reviewed fact-state (for the review watermark)
# --------------------------------------------------------------------------- #
# Justin's sheets are full WFIGS exports, so the acreage / containment / perimeter
# date HE WAS LOOKING AT are in the workbook itself. Those are the reviewed facts.
# We never substitute a later snapshot's numbers for them.
_FACT_COLUMNS = {
    "poly_GISAcres": "acres",
    "attr_IncidentSize": "incident_size",
    "attr_PercentContained": "pct_contained",
    "poly_DateCurrent": "perimeter_date_current",
}


def _num(v: Any) -> float | None:
    if v is None or (isinstance(v, str) and v.strip().lower() in ("", "null", "none", "nan")):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def sheet_facts(workbook: Path, sheet: str) -> dict[str, dict[str, Any]]:
    """{IRWIN ID -> reviewed facts} from one workbook sheet. Source values only.

    Unlike read_sheet() this keeps EVERY row, not just rows carrying a note: a sheet
    with no notes (Justin's 9-9) is still a completed review pass.
    """
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {}
    header = [str(h) if h is not None else "" for h in rows[0]]
    if "attr_IrwinID" not in header:
        raise ValueError(f"Sheet {sheet!r} lacks an attr_IrwinID column")
    i_irwin = header.index("attr_IrwinID")
    i_name = header.index("attr_IncidentName") if "attr_IncidentName" in header else None
    cols = {src: (header.index(src), dst) for src, dst in _FACT_COLUMNS.items() if src in header}

    out: dict[str, dict[str, Any]] = {}
    for r in rows[1:]:
        iid = r[i_irwin] if i_irwin < len(r) else None
        if not iid:
            continue
        facts: dict[str, Any] = {}
        for src, (i, dst) in cols.items():
            if i >= len(r):
                continue
            v = r[i]
            if dst == "perimeter_date_current":
                facts[dst] = v.isoformat(timespec="seconds") if hasattr(v, "isoformat") else (
                    str(v).strip() or None)
            else:
                facts[dst] = _num(v)
        # Same precedence as wfigs.load_wfigs_csv: perimeter GIS acres, else reported size.
        if facts.get("acres") is None and facts.get("incident_size") is not None:
            facts["acres"] = facts["incident_size"]
        facts.pop("incident_size", None)
        facts["fire_name"] = r[i_name] if i_name is not None and i_name < len(r) else None
        out[str(iid).strip().upper()] = facts
    return out


# "Moratorium until 9/21" / "moratorium until 9/21/26" - Justin's own recorded wording.
_UNTIL_RE = re.compile(r"\bmoratorium\s+until\s+(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?", re.I)


def expiry_from_note(note: str, review_year: int) -> str | None:
    """Deterministic parse of an expiry date already written in the analyst's note.

    Returns ISO 'YYYY-MM-DD' or None. The year is taken from the note when written,
    otherwise from the review year. This reads a date the analyst recorded; it does
    not read a bulletin PDF and does not infer an expiry that was never written down.
    """
    m = _UNTIL_RE.search(note or "")
    if not m:
        return None
    month, day, year = int(m.group(1)), int(m.group(2)), m.group(3)
    if year:
        y = int(year)
        y += 2000 if y < 100 else 0
    else:
        y = review_year
    try:
        return date(y, month, day).isoformat()
    except ValueError:
        return None


def seed_review_watermarks(records: dict[str, dict[str, Any]], workbook: Path = DEFAULT_WORKBOOK,
                           *, review_sheet: str = REVIEW_SHEET,
                           fallback_sheets: tuple[str, ...] = (DEFAULT_SHEET,),
                           by: str | None = None, force: bool = False) -> dict[str, Any]:
    """One-time backfill: treat `review_sheet` as the last completed review.

    The sheet is a full WFIGS export, so EVERY fire listed on it was part of the
    review - the rows Justin annotated are his conclusions, and the rows he left
    un-annotated are the ones he looked at and dismissed. Both get a watermark, so
    a dismissed fire stays quiet until its facts actually move.

    Fires the review sheet no longer lists (WFIGS drops contained fires) fall back
    to the most recent earlier sheet that does list them, and say so in
    facts_source. A fire with no facts on any sheet is left without a watermark
    rather than given invented numbers.
    """
    facts_by_sheet = {s: sheet_facts(workbook, s) for s in (review_sheet, *fallback_sheets)}
    review_year = _sheet_year(workbook, review_sheet)
    summary: dict[str, Any] = {"workbook": workbook.name, "review_sheet": review_sheet,
                               "seeded": 0, "seeded_from_fallback": 0, "skipped_existing": 0,
                               "records_created": 0, "no_facts": [], "moratorium_expiry_filled": []}

    # The reviewed universe: everything on the sheets, plus any record we already hold.
    reviewed_ids = {iid for f in facts_by_sheet.values() for iid in f} | set(records)
    for iid in sorted(reviewed_ids):
        if iid not in records:
            summary["records_created"] += 1
        stored = analyst.get(records, iid)
        if analyst.review_of(stored) and not force:
            summary["skipped_existing"] += 1
            continue
        sheet_used = next((s for s in (review_sheet, *fallback_sheets)
                           if iid in facts_by_sheet[s]), None)
        if sheet_used is None:
            summary["no_facts"].append(stored.get("fire_name") or iid)
            continue
        f = facts_by_sheet[sheet_used][iid]
        label = f"Justin workbook {review_sheet}"
        if sheet_used == review_sheet:
            source = (f"Justin workbook {review_sheet} (poly_GISAcres / attr_PercentContained / "
                      "poly_DateCurrent as reviewed)")
            summary["seeded"] += 1
        else:
            source = (f"Justin workbook {sheet_used} - not listed on {review_sheet} "
                      f"(dropped from the WFIGS feed by then)")
            summary["seeded_from_fallback"] += 1
        analyst.record_review(records, iid, stored.get("fire_name") or f.get("fire_name"),
                              facts=f, review_label=label, facts_source=source, by=by)

        # An expiry the analyst already wrote in their own note, made machine-readable.
        if stored.get("status") == "Existing Moratorium" and not (stored.get("moratorium_expires") or "").strip():
            iso = expiry_from_note(stored.get("notes") or "", review_year)
            if iso:
                records[iid.upper()]["moratorium_expires"] = iso
                records[iid.upper()]["moratorium_expiry_source"] = (
                    f"Analyst note, Justin workbook {DEFAULT_SHEET}")
                summary["moratorium_expiry_filled"].append(
                    {"fire": stored.get("fire_name") or iid, "expires": iso})
    return summary


def _sheet_year(workbook: Path, sheet: str) -> int:
    """Year for a M-D sheet name, from the workbook name ('2026-09 Fire Tracking')."""
    m = re.match(r"(\d{4})", workbook.name)
    return int(m.group(1)) if m else date.today().year


def import_baseline(records: dict[str, dict[str, Any]], workbook: Path = DEFAULT_WORKBOOK,
                    sheet: str = DEFAULT_SHEET) -> dict[str, Any]:
    """Merge baseline notes into `records` in place. Returns a summary dict."""
    rows = read_sheet(workbook, sheet)
    summary: dict[str, Any] = {"sheet": sheet, "rows_with_notes": len(rows), "imported": 0,
                               "skipped_existing": 0, "suggested": [], "workbook": workbook.name}
    source = f"{analyst.SOURCE_BASELINE_PREFIX} {sheet} (unconfirmed)"
    for row in rows:
        rec = analyst.get(records, row["irwin_id"], row["fire_name"])
        if (rec.get("notes") or "").strip() or rec["status"] != analyst.DEFAULT_STATUS:
            summary["skipped_existing"] += 1
            continue
        fields: dict[str, Any] = {"notes": f"[Justin workbook {sheet}] {row['note']}"}
        ref_bits = []
        if row["distance"] not in (None, "", "null ", "null"):
            ref_bits.append(f"Justin ref miles ({sheet}): {row['distance']}")
        if row["location"]:
            ref_bits.append(str(row["location"]).strip())
        fields["baseline_reference"] = "; ".join(ref_bits)
        suggested = suggest_status(row["note"])
        if suggested and suggested != analyst.DEFAULT_STATUS:
            fields["status"] = suggested
            summary["suggested"].append({"fire": row["fire_name"], "status": suggested})
        analyst.update(records, row["irwin_id"], row["fire_name"], source=source, **fields)
        summary["imported"] += 1
    return summary
